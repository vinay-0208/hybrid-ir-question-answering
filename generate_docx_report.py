"""
Script to generate the comprehensive Word Implementation Report
in reports/Implementation_Report_23MID0429.docx adhering to academic standards
and reflecting the corrected, verified evaluation metrics.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def create_report():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    reports_dir = os.path.join(base_dir, "reports")
    docx_path = os.path.join(reports_dir, "Implementation_Report_23MID0429.docx")

    doc = docx.Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # ---------------------------------------------------------
    # 1. COVER PAGE
    # ---------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("VELLORE INSTITUTE OF TECHNOLOGY\nSCHOOL OF COMPUTER SCIENCE AND ENGINEERING")
    r_inst.bold = True
    r_inst.font.size = Pt(14)
    r_inst.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = p_course.add_run("CSI4001\nNATURAL LANGUAGE PROCESSING AND COMPUTATIONAL LINGUISTICS\n")
    r_c1.bold = True
    r_c1.font.size = Pt(13)
    r_c1.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    r_c2 = p_course.add_run("CASE STUDY 2")
    r_c2.bold = True
    r_c2.font.size = Pt(12)
    r_c2.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph().paragraph_format.space_after = Pt(28)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("IMPLEMENTATION REPORT\n\nHybrid Transformer-Based Information Retrieval\nand Question Answering")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph().paragraph_format.space_after = Pt(36)

    # Student Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    data = [
        ("Student Name", "Vinay Vishwanathan"),
        ("Register Number", "23MID0429"),
        ("Course Code & Title", "CSI4001: Natural Language Processing and Computational Linguistics"),
        ("Case Study & Topic", "Case Study 2: Hybrid Transformer-Based Information Retrieval and Question Answering"),
        ("Submission Date", "October 2026")
    ]

    for i, (k, v) in enumerate(data):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.5)
        c1.width = Inches(3.8)

        set_cell_background(c0, "F1F5F9")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_page_break()

    # Helper function for headings
    def add_sec_heading(title, level=1):
        h = doc.add_heading(title, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run = h.runs[0]
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif level == 2:
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        return h

    # ---------------------------------------------------------
    # 2. IMPLEMENTATION OVERVIEW
    # ---------------------------------------------------------
    add_sec_heading("1. Implementation Overview")
    doc.add_paragraph(
        "Modern information retrieval (IR) and question answering (QA) systems must satisfy dual demands: "
        "high recall across large unstructured document collections and high precision when extracting direct, verifiable factual answers. "
        "Traditional lexical keyword retrieval based on the Okapi BM25 ranking algorithm is fast and exact for specific terminology, but fails "
        "when queries contain synonyms, rephrasings, or conceptual abstractions (the vocabulary mismatch problem). "
        "Conversely, dense semantic retrieval using Transformer bi-encoders maps queries and passages into continuous geometric vector spaces, "
        "capturing semantic depth but occasionally overlooking exact keyword specifications."
    )
    doc.add_paragraph(
        "This project implements a fully functional, production-grade hybrid retrieval and question answering system. "
        "The system combines Okapi BM25 lexical search with Sentence-BERT dense semantic search (sentence-transformers/all-MiniLM-L6-v2), "
        "fuses their ranking candidates via Reciprocal Rank Fusion (RRF), reorders the top candidate passages using a Transformer cross-encoder "
        "(cross-encoder/ms-marco-MiniLM-L-6-v2), extracts grounded factual answers using a local SQuAD-trained DistilBERT extractive QA model, "
        "and enforces strict abstention policies with provenance tracking to eliminate ungrounded hallucinations."
    )

    # ---------------------------------------------------------
    # 3. TECHNOLOGIES AND CONFIGURATION
    # ---------------------------------------------------------
    add_sec_heading("2. Technologies and Configuration")
    doc.add_paragraph(
        "The software stack was selected to guarantee 100% local execution without requiring paid cloud APIs, "
        "external API keys, or cloud LLM subscriptions:"
    )

    t_tech = doc.add_table(rows=7, cols=3)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_headers = ["Pipeline Component", "Technology / Model", "Role & Configuration Justification"]
    for j, h in enumerate(tech_headers):
        cell = t_tech.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tech_rows = [
        ("Web Application UI", "Streamlit 1.65.0", "Interactive dashboard with Heuristic ReAct trace, HyDE preview, 5-way comparison, and live RAGAS proxy metrics."),
        ("Heuristic Query Rewriting & HyDE", "Controlled 2-Step Heuristic ReAct + HyDE (alpha=0.6)", "Deterministic rule-based query expansion with 2-step RRF fusion and corpus-snippet PRF HyDE synthesis."),
        ("Lexical Retrieval", "rank-bm25 0.2.2 & Custom Fallback", "Okapi BM25 with k1=1.5, b=0.75 for fast, exact-match keyword candidate retrieval."),
        ("Dense Semantic & ANN Search", "all-MiniLM-L6-v2 + FAISS 1.15.1 (IndexHNSWFlat)", "384-d L2-normalized embeddings indexed in FAISS HNSW graph (METRIC_INNER_PRODUCT) with NumPy fallback."),
        ("Ranking Fusion & Reranking", "RRF (k=60) + cross-encoder/ms-marco-MiniLM-L-6-v2", "Unsupervised rank aggregation followed by full cross-attention precision reranking."),
        ("QA, Extractive Synthesis & Evaluation", "distilbert-base-cased-distilled-squad + Extractive Synthesizer", "Extractive span QA, multi-passage extractive sentence synthesis with inline citations [1: pid], and RAGAS-inspired proxy evaluation.")
    ]

    for i, (c1, c2, c3) in enumerate(tech_rows, start=1):
        row = t_tech.rows[i]
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate([c1, c2, c3]):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0:
                r.bold = True

    # ---------------------------------------------------------
    # 4. APPLICATION ARCHITECTURE AND DIAGRAM
    # ---------------------------------------------------------
    add_sec_heading("3. Application Architecture and Pipeline Flow")
    doc.add_paragraph(
        "The system architecture is partitioned into two clear stages: an Offline Indexing Stage and an Online Query Execution Stage."
    )
    doc.add_paragraph(
        "1. Offline Indexing Stage: Document ingestion extracts raw text from .txt and .pdf files, normalizes whitespace, "
        "and segments documents into sliding-window passages (default: 150 words per passage with 30-word overlap). "
        "Passages are indexed into an Okapi BM25 inverted index while concurrently being encoded into 384-dimensional L2-normalized dense vectors "
        "using all-MiniLM-L6-v2 and indexed into a FAISS Approximate Nearest Neighbor (IndexHNSWFlat) graph.\n"
        "2. Online Query Execution Stage: In Mode 5, a user query passes through a bounded 2-step Heuristic ReAct Query Rewriter (Thought -> Action -> Observation) "
        "that fuses original and expanded lexical candidates via weighted RRF, alongside a Hypothetical Document Embeddings (HyDE) generator using top-1 corpus pseudo-relevance feedback. "
        "Parallel BM25 and FAISS HNSW Dense retrieval fetch top candidate pools, which are harmonized via Reciprocal Rank Fusion (RRF, k=60) and precision-reranked using the MS MARCO cross-encoder. "
        "Finally, DistilBERT extractive QA and the Multi-Passage Extractive Answer Synthesizer assemble a verified response from corpus sentences with inline source citations [1: passage_id] and compute RAGAS-inspired heuristic proxy metrics."
    )

    diag_path = os.path.join(reports_dir, "architecture_diagram.png")
    if os.path.exists(diag_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(diag_path, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: Complete Offline and Online Pipeline Architecture for Hybrid IR and QA.")
        r_cap.font.size = Pt(9.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ---------------------------------------------------------
    # 5. DOCUMENT INGESTION AND INDEXING
    # ---------------------------------------------------------
    add_sec_heading("4. Document Ingestion and Passage Segmentation")
    doc.add_paragraph(
        "The ingestion module (app/ingestion.py) handles heterogeneous file types (.txt and .pdf) with robust exception handling. "
        "Empty files, unprintable characters, and encoding mismatches (UTF-8 with Latin-1 fallback) are cleanly intercepted. "
        "Passage segmentation employs a configurable sliding window algorithm: given word count W, chunk size S, and overlap O, "
        "the step size is step = S - O. Each passage is assigned a unique passage identifier (e.g., 01_information_retrieval_fundamentals.txt_p0), "
        "preserving full document provenance."
    )

    s3_path = os.path.join(reports_dir, "screenshot_3_ingestion.png")
    if os.path.exists(s3_path):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(s3_path, width=Inches(6.2))
        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap3 = p_cap3.add_run("Figure 2: Live Ingestion and Corpus Management Interface showing loaded documents and passage segmentation.")
        r_cap3.font.size = Pt(9.5)
        r_cap3.italic = True
        r_cap3.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ---------------------------------------------------------
    # 6. RETRIEVAL AND QUESTION ANSWERING DEMONSTRATION
    # ---------------------------------------------------------
    add_sec_heading("5. Retrieval, Question Answering, and Evidence Grounding")
    doc.add_paragraph(
        "In the primary query interface (app/streamlit_app.py Tab 1), users enter natural-language questions and select the execution pipeline mode. "
        "The application returns a structured answer card featuring:\n"
        "• The multi-passage extractive synthesized answer (combining verified corpus sentences with inline citations [1: passage_id]) alongside the concise DistilBERT span.\n"
        "• The Heuristic ReAct Query Rewriting trace (Thought -> Action -> Observation) and pseudo-relevance feedback HyDE hypothetical document.\n"
        "• Live RAGAS-inspired heuristic proxy scores (Faithfulness Proxy, Answer Relevance Proxy, Context Precision, Context Recall).\n"
        "• An operational latency breakdown detailing retrieval, fusion, reranking, and QA execution time in milliseconds.\n"
        "• An ordinal ranking table showing candidate ranks, BM25 scores, FAISS Dense cosine scores, RRF scores, and cross-encoder scores.\n"
        "• An explicit abstention policy: when out-of-domain queries (such as 'What is the capital of planet Mars?') are submitted, "
        "the system rejects ungrounded generation across all five modes and outputs: 'Insufficient evidence in the knowledge base to answer this question reliably.'"
    )

    s1_path = os.path.join(reports_dir, "screenshot_1_home_query.png")
    if os.path.exists(s1_path):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(s1_path, width=Inches(6.2))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap1 = p_cap1.add_run("Figure 3: Live Query Execution showing Extracted Answer, Supporting Evidence, and Multi-Stage Ranks.")
        r_cap1.font.size = Pt(9.5)
        r_cap1.italic = True
        r_cap1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ---------------------------------------------------------
    # 7. EVALUATION METRICS AND ACTUAL RESULTS
    # ---------------------------------------------------------
    add_sec_heading("6. Evaluation Metrics and Actual Experimental Results")
    doc.add_paragraph(
        "The evaluation module (app/metrics.py) implements standard reference-based IR/QA metrics alongside local RAGAS-inspired heuristic proxy metrics. "
        "All metrics are strictly normalized and bounded in the range [0.0, 1.0]:\n"
        "1. Precision@k: P@k = |Unique relevant IDs in top-k| / k.\n"
        "2. Recall@k: R@k = |Unique relevant IDs in top-k| / |Total relevant IDs|.\n"
        "3. Mean Reciprocal Rank (MRR): Reciprocal rank 1/r of the first relevant passage retrieved.\n"
        "4. Normalized Discounted Cumulative Gain (NDCG@k): DCG@k normalized by Ideal DCG (IDCG@k) without duplicate gain inflation.\n"
        "5. Exact Match (EM) & Token-Level F1: Normalized lexical match and harmonic mean of token precision and recall.\n"
        "6. RAGAS-Inspired Heuristic Proxies (Faithfulness, Answer Relevance, Context Precision, Context Recall): Deterministic local approximations evaluating sentence-level lexical/negation grounding in retrieved passages, bi-encoder cosine + lexical query-answer alignment, average precision of relevant context ranks, and reference token/sentence coverage."
    )
    doc.add_paragraph(
        "The automated benchmark runner evaluated 10 ground-truth query-answer pairs across all five retrieval modalities on the indexed knowledge base (after model warmup). "
        "The verified empirical results recorded on this machine are presented below:"
    )

    # Summary table from benchmark
    t_res = doc.add_table(rows=6, cols=11)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    eval_cols = ["Retrieval Mode", "P@5", "R@5", "MRR", "NDCG@5", "EM", "Token F1", "Faithful", "Ans Rel", "Ctx Prec", "Latency"]
    for j, h in enumerate(eval_cols):
        cell = t_res.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=60, bottom=60, left=45, right=45)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    eval_data = [
        ("1. BM25 Only", "0.180", "0.900", "0.850", "0.863", "0.200", "0.433", "1.000", "0.674", "0.850", "222.07 ms"),
        ("2. Dense Only (FAISS)", "0.180", "0.900", "0.770", "0.802", "0.200", "0.373", "1.000", "0.629", "0.770", "224.09 ms"),
        ("3. Hybrid (RRF)", "0.160", "0.800", "0.750", "0.763", "0.200", "0.373", "1.000", "0.662", "0.750", "239.55 ms"),
        ("4. Hybrid + Cross-Encoder", "0.200", "1.000", "0.875", "0.906", "0.200", "0.383", "1.000", "0.667", "0.875", "627.86 ms"),
        ("5. Heuristic ReAct + HyDE + Rerank", "0.200", "1.000", "0.950", "0.963", "0.300", "0.543", "1.000", "0.658", "0.950", "674.65 ms")
    ]

    for i, row_vals in enumerate(eval_data, start=1):
        row = t_res.rows[i]
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row_vals):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=50, bottom=50, left=45, right=45)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(7.5)
            if j == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    s4_path = os.path.join(reports_dir, "screenshot_4_evaluation.png")
    if os.path.exists(s4_path):
        p_img4 = doc.add_paragraph()
        p_img4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(s4_path, width=Inches(6.2))
        p_cap4 = doc.add_paragraph()
        p_cap4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap4 = p_cap4.add_run("Figure 4: Automated Evaluation Dashboard showing corrected empirical metrics and visual bar charts.")
        r_cap4.font.size = Pt(9.5)
        r_cap4.italic = True
        r_cap4.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ---------------------------------------------------------
    # 8. COMPARISON OF RETRIEVAL METHODS
    # ---------------------------------------------------------
    add_sec_heading("7. Comparative Analysis of Retrieval Methods")
    doc.add_paragraph(
        "Analysis of the empirical benchmark results across all 5 modes yields clear academic and engineering conclusions:\n"
        "• Impact of Heuristic ReAct Query Rewriting & HyDE (Mode 5): Mode 5 (Heuristic ReAct + HyDE + Rerank) achieved Recall@5 = 1.000, "
        "MRR = 0.950, NDCG@5 = 0.963, Context Precision = 0.950, Exact Match = 0.300, and Token F1 = 0.543. Compared to Mode 4 (MRR = 0.875, NDCG@5 = 0.906, EM = 0.200, Token F1 = 0.383), "
        "the 2-step heuristic query expansion and pseudo-relevance feedback HyDE vector blending promoted target passages on conceptual/parameter queries (q2 and q5) to Rank 1, "
        "while placing 9 of 10 target passages at Rank 1 and q6 at Rank 2.\n"
        "• Baseline Hybrid + Cross-Encoder (Mode 4): Achieved Recall@5 = 1.000, Precision@5 = 0.200 (the theoretical ceiling for 1 target passage per query at k=5), "
        "MRR = 0.875, and NDCG@5 = 0.906, outperforming un-reranked BM25 (0.863), Dense FAISS (0.802), and Hybrid RRF (0.763).\n"
        "• Extractive Grounding & Faithfulness Proxy: All 5 modes achieved a Faithfulness proxy score of 1.000 on answerable benchmark queries because the Grounded Answer Synthesizer "
        "is a verbatim multi-passage extractive sentence selector with inline citations [1: passage_id], ensuring every synthesized sentence comes directly from the retrieved context.\n"
        "• Warmup-Corrected Latency Trade-Off: First-stage retrieval modes (BM25, Dense FAISS, and Hybrid RRF) ran in ~222–240 ms per query (dominated by DistilBERT reader inference), "
        "whereas Cross-Encoder modes required ~628 ms (Mode 4) and ~675 ms (Mode 5) due to joint query-passage Transformer scoring."
    )

    s2_path = os.path.join(reports_dir, "screenshot_2_comparison.png")
    if os.path.exists(s2_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(s2_path, width=Inches(6.2))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = p_cap2.add_run("Figure 5: Side-by-Side Comparison interface demonstrating ranking and latency variations.")
        r_cap2.font.size = Pt(9.5)
        r_cap2.italic = True
        r_cap2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ---------------------------------------------------------
    # 9. AUTOMATED TESTING
    # ---------------------------------------------------------
    add_sec_heading("8. Automated Testing and Quality Verification")
    doc.add_paragraph(
        "The automated test suite in tests/ comprises 49 unit and integration tests across seven dedicated test modules:\n"
        "• test_query_transform.py: Validated Heuristic ReAct query analysis, non-leaky concept expansion, empty-input handling, 2-step RRF iterative retrieval, and HyDE synthesis (5 tests).\n"
        "• test_metrics.py: Validated Exact Match, Token F1, Precision@k, Recall@k, MRR, NDCG@k, duplicate candidate handling, benchmark label consistency, and RAGAS-inspired Faithfulness (including negation/fabrication penalties), Answer Relevance, Context Precision, and Context Recall (13 tests).\n"
        "• test_ingestion.py: Validated text cleaning, UTF-8/Latin-1 decoding, sliding-window passage chunking, and repository lifecycle (4 tests).\n"
        "• test_retrieval.py: Validated query tokenization, Okapi BM25 scoring, pure-Python BM25 fallback, RRF, Convex Combination, FAISS IndexHNSWFlat ANN search, HyDE vector blending, and FAISS vs. exact NumPy consistency (7 tests).\n"
        "• test_reranking.py: Validated CrossEncoder initialization, candidate reordering, and fallback mechanisms (2 tests).\n"
        "• test_answering.py: Validated sentence splitting, evidence selection, abstention policies, heuristic fallback, and GroundedAnswerSynthesizer inline citations and ungrounded passage rejection (7 tests).\n"
        "• test_integration.py: Validated end-to-end engine execution across all 5 modes, baseline restoration when rewriting/HyDE are disabled, out-of-domain abstention across all 5 modes, forced offline fallback mode, and JSON/CSV serialization (11 tests).\n\n"
        "Test Execution Result: All 49 tests passed successfully (0 failures, 0 errors, 0 skipped)."
    )

    # ---------------------------------------------------------
    # 10. LIMITATIONS AND FUTURE IMPROVEMENTS
    # ---------------------------------------------------------
    add_sec_heading("9. Limitations and Future Improvements")
    doc.add_paragraph(
        "1. Heuristic ReAct/HyDE & Extractive Synthesis vs. Generative LLM: Because this machine operates on CPU (CUDA=False) with ~1.66 GB of available physical RAM, "
        "the query rewriter, HyDE generator, and grounded answer synthesizer are implemented as deterministic heuristic and multi-passage extractive components rather than generative causal LLMs. "
        "Similarly, the RAGAS evaluation scores are deterministic local heuristic proxies rather than official LLM-as-a-judge RAGAS calls.\n"
        "2. FAISS ANN Scale & Single-Target Precision@5 Ceiling: On the 13-passage built-in corpus, FAISS IndexHNSWFlat provides architectural completeness and exact consistency with NumPy dot-product search rather than a measurable ANN speed advantage (which emerges at >10,000 passages). Furthermore, with 1 target passage per query, Precision@5 has a mathematical upper bound of 1/5 = 0.200.\n"
        "3. Cross-Encoder Latency: Cross-encoder reranking latency (~628–675 ms on CPU) can be further accelerated via ONNX Runtime INT8 quantization."
    )

    # ---------------------------------------------------------
    # 11. SOURCE CODE AND EXECUTION INSTRUCTIONS
    # ---------------------------------------------------------
    add_sec_heading("10. Source Code and Execution Instructions")
    doc.add_paragraph(
        "Execution on Windows:\n"
        "1. Navigate to the project root directory: cd \"c:\\Fall Semister\\2026\\nlp\\case study-2\"\n"
        "2. Run the automated launcher: run_app.bat\n"
        "   Or manually launch Streamlit:\n"
        "   .\\.venv\\Scripts\\streamlit.exe run app/streamlit_app.py\n\n"
        "Running the Test Suite:\n"
        "   .\\.venv\\Scripts\\pytest.exe tests/ -v\n\n"
        "Running Benchmark Evaluation:\n"
        "   .\\.venv\\Scripts\\python.exe run_benchmark_eval.py\n\n"
        "Artifact File Locations:\n"
        "• Implementation Report: reports/Implementation_Report_23MID0429.docx\n"
        "• Architecture Diagram: reports/architecture_diagram.png\n"
        "• Benchmark Summary: reports/benchmark_summary.csv and reports/benchmark_summary.json\n"
        "• Captured UI Screenshots: reports/screenshot_1_home_query.png through screenshot_4_evaluation.png"
    )

    # Save document
    doc.save(docx_path)
    print(f"Implementation report successfully generated at: {docx_path}")

if __name__ == "__main__":
    create_report()
